import io
import re
import time
from typing import Any, Dict, List, Optional, Tuple

import docx
import pdfplumber
import pypdf

from backend.app.models.profile import (
    EducationItem,
    ExperienceItem,
    ParsedResume,
    ProjectItem,
)
from backend.app.services.parser.normalizer import (
    GITHUB_REGEX,
    detect_section_header,
    extract_contact_info,
    extract_metrics,
    extract_skills_from_text,
    normalize_text,
)


class ParsingError(Exception):
    """Base exception for document parsing failures."""
    pass


class ScannedDocumentError(ParsingError):
    """Raised when an uploaded PDF has no extractable digital text (scanned image)."""
    pass


class EmptyDocumentError(ParsingError):
    """Raised when an uploaded file contains no readable content."""
    pass


class UnsupportedFormatError(ParsingError):
    """Raised when the file format is not supported."""
    pass


class ResumeParser:
    """
    Deterministic Resume Parsing Service.
    Extracts text from PDF, DOCX, and TXT files, performs section segmentation,
    and structures candidate information into a validated ParsedResume object.
    """

    SUPPORTED_EXTENSIONS = [".pdf", ".docx", ".txt"]

    @classmethod
    def parse_bytes(cls, file_bytes: bytes, filename: str) -> ParsedResume:
        """
        Parse raw bytes of an uploaded resume file into a structured ParsedResume object.
        """
        start_time = time.perf_counter()

        if not file_bytes or len(file_bytes.strip()) == 0:
            raise EmptyDocumentError("The uploaded file is empty.")

        ext = "." + filename.lower().split(".")[-1] if "." in filename else ""
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise UnsupportedFormatError(
                f"Unsupported file format '{ext}'. Allowed formats: {', '.join(cls.SUPPORTED_EXTENSIONS)}"
            )

        raw_text, page_count = cls._extract_raw_text(file_bytes, ext)
        sanitized_text = normalize_text(raw_text)

        if len(sanitized_text.strip()) < 20:
            if ext == ".pdf":
                raise ScannedDocumentError(
                    "The uploaded PDF contains no extractable text or appears to be a scanned image. "
                    "Please upload a text-based PDF."
                )
            raise EmptyDocumentError("The document does not contain sufficient text for parsing.")

        parsed_resume = cls._parse_structured_text(sanitized_text)

        elapsed_ms = int((time.perf_counter() - start_time) * 1000)
        parsed_resume.extraction_metadata = {
            "filename": filename,
            "file_type": ext.lstrip("."),
            "pages_parsed": page_count,
            "character_count": len(sanitized_text),
            "parsing_time_ms": elapsed_ms,
        }
        parsed_resume.raw_text = sanitized_text

        return parsed_resume

    @classmethod
    def _extract_raw_text(cls, file_bytes: bytes, ext: str) -> Tuple[str, int]:
        """Extract text from file bytes based on format."""
        if ext == ".txt":
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="ignore")
            return text, 1

        if ext == ".docx":
            doc_file = io.BytesIO(file_bytes)
            try:
                doc = docx.Document(doc_file)
                full_text = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        full_text.append(para.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            full_text.append(" | ".join(row_text))
                return "\n".join(full_text), 1
            except Exception as e:
                raise ParsingError(f"Failed to read DOCX file: {e}")

        if ext == ".pdf":
            pages_text = []
            page_count = 0
            # Primary extraction with pdfplumber
            try:
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    page_count = len(pdf.pages)
                    for page in pdf.pages:
                        extracted = page.extract_text(layout=True) or page.extract_text()
                        if extracted:
                            pages_text.append(extracted)
            except Exception:
                # Fallback to pypdf if pdfplumber fails
                try:
                    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                    page_count = len(reader.pages)
                    for page in reader.pages:
                        extracted = page.extract_text()
                        if extracted:
                            pages_text.append(extracted)
                except Exception as e:
                    raise ParsingError(f"Failed to read PDF file: {e}")

            return "\n".join(pages_text), page_count

        raise UnsupportedFormatError(f"Cannot process format: {ext}")

    @classmethod
    def _parse_structured_text(cls, text: str) -> ParsedResume:
        """Segment raw text into sections and construct structured ParsedResume."""
        lines = [line.strip() for line in text.splitlines() if line.strip()]

        # 1. Candidate Name Extraction
        candidate_name = cls._extract_candidate_name(lines)

        # 2. Contact Information
        contacts = extract_contact_info(text)

        # 3. Section Segmentation
        sections = cls._segment_sections(lines)

        # 4. Extract Projects
        projects = cls._extract_projects(sections.get("PROJECTS", []))

        # 5. Extract Education
        education = cls._extract_education(sections.get("EDUCATION", []))

        # 6. Extract Experience
        experiences = cls._extract_experience(sections.get("EXPERIENCE", []))

        # 7. Extract Skills (from dedicated SKILLS section + entire resume)
        skills_section_text = "\n".join(sections.get("SKILLS", []))
        section_skills, _ = extract_skills_from_text(skills_section_text)
        overall_skills, _ = extract_skills_from_text(text)
        all_skills = sorted(list(set(section_skills + overall_skills)))

        return ParsedResume(
            full_name=candidate_name,
            email=contacts["email"],
            phone=contacts["phone"],
            github_url=contacts["github_url"],
            linkedin_url=contacts["linkedin_url"],
            education=education,
            experiences=experiences,
            projects=projects,
            skills=all_skills,
        )

    @classmethod
    def _extract_candidate_name(cls, lines: List[str]) -> str:
        """Infer candidate name from top lines of the resume."""
        for line in lines[:5]:
            # Skip lines that look like emails, phone numbers, or URLs
            if "@" in line or "http" in line or "www." in line or "/" in line:
                continue
            if re.search(r"\d{5,}", line):
                continue
            # Remove symbols
            cleaned = re.sub(r"[|•,-]", " ", line).strip()
            words = cleaned.split()
            if 1 <= len(words) <= 4 and all(w.replace(".", "").isalpha() for w in words):
                return cleaned.title()

        return "Candidate"

    @classmethod
    def _segment_sections(cls, lines: List[str]) -> Dict[str, List[str]]:
        """Split document lines into categorized section blocks."""
        sections: Dict[str, List[str]] = {
            "HEADER": [],
            "EDUCATION": [],
            "EXPERIENCE": [],
            "PROJECTS": [],
            "SKILLS": [],
            "CERTIFICATIONS": [],
        }

        current_section = "HEADER"

        for line in lines:
            detected = detect_section_header(line)
            if detected:
                current_section = detected
                continue
            sections[current_section].append(line)

        return sections

    @classmethod
    def _extract_projects(cls, lines: List[str]) -> List[ProjectItem]:
        """Extract individual projects from project section lines."""
        if not lines:
            return []

        projects: List[ProjectItem] = []
        current_title: Optional[str] = None
        current_repo: Optional[str] = None
        current_bullets: List[str] = []

        for line in lines:
            line_clean = line.strip()
            if not line_clean:
                continue

            # Check if line is a standalone URL
            gh_match = GITHUB_REGEX.search(line_clean)
            repo_url = gh_match.group(0) if gh_match else None
            is_url_line = line_clean.startswith("http://") or line_clean.startswith("https://")

            if is_url_line:
                if repo_url and current_title:
                    current_repo = repo_url
                continue

            # Check if this line is a bullet point
            is_bullet = line_clean.startswith("-") or line_clean.startswith("•") or line_clean.startswith("*")

            if not is_bullet and len(line_clean) < 80 and not line_clean.endswith("."):
                # Potential new project title
                if current_title:
                    bullets_text = " ".join(current_bullets)
                    tech_stack, _ = extract_skills_from_text(f"{current_title} {bullets_text}")
                    projects.append(
                        ProjectItem(
                            title=current_title,
                            repo_url=current_repo,
                            tech_stack=tech_stack,
                            bullets=current_bullets,
                            metrics_detected=extract_metrics(bullets_text),
                        )
                    )
                    current_bullets = []
                    current_repo = None

                # Clean title: strip URLs if present in the title line
                title_clean = re.sub(r"https?://\S+", "", line_clean).strip().strip("|-()")
                current_title = title_clean or "Project"
                if repo_url:
                    current_repo = repo_url
            else:
                clean_bullet = line_clean.lstrip("-•* ").strip()
                if clean_bullet:
                    current_bullets.append(clean_bullet)
                if repo_url and not current_repo:
                    current_repo = repo_url

        # Append last pending project
        if current_title:
            bullets_text = " ".join(current_bullets)
            tech_stack, _ = extract_skills_from_text(f"{current_title} {bullets_text}")
            projects.append(
                ProjectItem(
                    title=current_title,
                    repo_url=current_repo,
                    tech_stack=tech_stack,
                    bullets=current_bullets,
                    metrics_detected=extract_metrics(bullets_text),
                )
            )

        return projects

    @classmethod
    def _extract_education(cls, lines: List[str]) -> List[EducationItem]:
        """Extract education records from education section lines."""
        if not lines:
            return []

        education_items: List[EducationItem] = []
        combined_text = " \n ".join(lines)

        # Look for years (e.g. 2022 - 2026 or 2026)
        years = [int(y) for y in re.findall(r"\b(20\d{2}|19\d{2})\b", combined_text)]
        grad_year = max(years) if years else None
        start_year = min(years) if len(years) > 1 else None

        # Look for GPA
        gpa_match = re.search(r"\b(?:gpa|cgpa|percentage)?\s*[:=]?\s*([0-9]\.[0-9]{1,2}|10\.0|[0-9]{2}(?:\.[0-9])?%?)\b", combined_text, re.IGNORECASE)
        gpa = None
        if gpa_match:
            try:
                raw_gpa = gpa_match.group(1).rstrip("%")
                gpa = float(raw_gpa)
            except ValueError:
                pass

        # Identify degree and institution
        degree = "Bachelor of Technology"
        for line in lines:
            lower = line.lower()
            if any(d in lower for d in ["b.tech", "btech", "b.e.", "bachelor", "b.s.", "m.tech", "master", "m.s."]):
                degree = line.strip("|-• ")
                break

        institution = lines[0].strip("|-• ") if lines else "University"

        education_items.append(
            EducationItem(
                institution=institution,
                degree=degree,
                gpa=gpa,
                start_year=start_year,
                grad_year=grad_year,
            )
        )
        return education_items

    @classmethod
    def _extract_experience(cls, lines: List[str]) -> List[ExperienceItem]:
        """Extract experience entries from experience section lines."""
        if not lines:
            return []

        experiences: List[ExperienceItem] = []
        current_role: Optional[str] = None
        current_company: Optional[str] = None
        current_duration: Optional[str] = None
        current_bullets: List[str] = []

        ROLE_KEYWORDS = [
            "intern", "engineer", "developer", "analyst", "manager", "lead",
            "architect", "consultant", "specialist", "designer", "associate",
            "officer", "assistant"
        ]

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            is_bullet = line_str.startswith("-") or line_str.startswith("•") or line_str.startswith("*")

            # Look for date patterns in line (e.g. May 2025 - Aug 2025)
            date_match = re.search(
                r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s*\d{4}\s*(?:-|–|to)\s*(?:Present|Current|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s*\d{4})\b",
                line_str,
                re.IGNORECASE,
            )

            # If line is solely or primarily a date line and we have an active role/company
            if date_match and len(line_str) < 40 and not is_bullet and current_company:
                current_duration = date_match.group(0)
                continue

            if not is_bullet and len(line_str) < 70 and not line_str.endswith("."):
                if current_role and current_company:
                    experiences.append(
                        ExperienceItem(
                            company=current_company,
                            role=current_role,
                            duration=current_duration,
                            bullets=current_bullets,
                        )
                    )
                    current_bullets = []
                    current_duration = None

                # Split role and company if separator present
                parts = re.split(r"\s*[|–—@-]\s*", line_str)
                if len(parts) >= 2:
                    part_a = parts[0].strip()
                    part_b = parts[1].strip()
                    if any(k in part_b.lower() for k in ROLE_KEYWORDS) and not any(k in part_a.lower() for k in ROLE_KEYWORDS):
                        current_company = part_a
                        current_role = part_b
                    elif any(k in part_a.lower() for k in ROLE_KEYWORDS) and not any(k in part_b.lower() for k in ROLE_KEYWORDS):
                        current_role = part_a
                        current_company = part_b
                    else:
                        current_company = part_a
                        current_role = part_b
                else:
                    if any(k in line_str.lower() for k in ROLE_KEYWORDS):
                        current_role = line_str
                        current_company = "Organization"
                    else:
                        current_company = line_str
                        current_role = "Role"

                if date_match and not current_duration:
                    current_duration = date_match.group(0)
            else:
                clean_bullet = line_str.lstrip("-•* ").strip()
                if clean_bullet:
                    current_bullets.append(clean_bullet)

        if current_role and current_company:
            experiences.append(
                ExperienceItem(
                    company=current_company,
                    role=current_role,
                    duration=current_duration,
                    bullets=current_bullets,
                )
            )

        return experiences
